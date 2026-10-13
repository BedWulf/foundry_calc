"""Тесты сводной ведомости (core.summary): сборка и экспорт в Excel."""
from __future__ import annotations

from pathlib import Path
from typing import Dict, Any

import pytest

from core.project import Project
from core.summary import (
    SummaryResult, SummaryRow, build_summary, export_to_excel,
)


@pytest.fixture()
def full_summaries() -> Dict[str, Dict[str, Any]]:
    """Сводки всех отделений — как их отдаёт MainWindow.department_summaries."""
    return {
        "melting": {"metal_charge_t": 58932.0, "furnace_name": "ИСТ-6М1",
                    "furnace_count": 6, "load_factor": 0.842,
                    "material_category": "steel"},
        "molding": {"annual_molds_total": 120000, "line_name": "ИЛ-225",
                    "line_count": 2, "load_factor": 0.78,
                    "material_category": "steel", "shakeout_in_line": True},
        "cores": {"annual_cores": 240000, "machine_name": "СБ-834",
                  "machine_count": 4, "load_factor": 0.71},
        "sand_prep": {"total_park_units": 9, "avg_load_factor": 0.75,
                      "park": {"mixer": {"title": "Смесительный штыковой",
                                         "model": "СБ-164", "units": 3,
                                         "load_factor": 0.72}}},
        "cleaning": {"total_units": 12, "avg_load_factor": 0.8,
                     "park": {"shot": {"title": "Дробемётная камера",
                                        "machine_count": 2,
                                        "load_factor": 0.66}}},
        "warehouses": {"total_area_m2": 1250.0,
                       "park": {"metal": {"title": "Склад металла",
                                          "unit": "Навес 24х12",
                                          "count": 2}}},
        "workforce": {"total": 212, "main_total": 66, "aux_total": 146},
        "areas": {"building_total_m2": 6800.0},
        "transport": {"cranes_total": 7,
                      "rows": [{"name": "Кран мостовой 20/5 т",
                                "units": 3, "department": "Плавильное"}]},
    }


class TestBuildSummary:
    """Сборка сводки из сводок отделений."""

    def test_kpis(self, full_summaries) -> None:
        proj = Project(annual_output_t=42500.0, name="Цех №1")
        res = build_summary(full_summaries, proj)
        assert res.project_name == "Цех №1"
        assert res.annual_output_t == 42500.0
        assert res.material_category == "steel"
        assert res.metal_charge_t == pytest.approx(58932.0)
        assert res.annual_molds == 120000
        assert res.annual_cores == 240000
        assert res.workforce_total == 212
        assert res.building_area_m2 == pytest.approx(6800.0)
        assert res.cranes_total == 7

    def test_rows_cover_departments(self, full_summaries) -> None:
        res = build_summary(full_summaries)
        depts = {r.department for r in res.rows}
        assert {"Плавильное", "Формовочное", "Стержневое",
                "Смесеприготовительное", "Обрубка/очистка",
                "Склады", "Транспорт"} <= depts
        melting = next(r for r in res.rows if r.department == "Плавильное")
        assert melting.equipment == "ИСТ-6М1"
        assert melting.units == 6
        assert melting.load_factor == pytest.approx(0.842)

    def test_empty_summaries_marks_missing(self) -> None:
        res = build_summary({})
        assert res.rows == []
        assert len(res.missing_sections) == 9
        assert "Плавильное отделение" in res.missing_sections

    def test_partial_summaries(self) -> None:
        res = build_summary({"melting": {"furnace_name": "ИСТ-6М1",
                                         "furnace_count": 6,
                                         "load_factor": 0.84,
                                         "metal_charge_t": 58932.0}})
        assert len(res.rows) == 1
        assert res.metal_charge_t == pytest.approx(58932.0)
        # 8 разделов без данных
        assert len(res.missing_sections) == 8

    def test_as_dict_roundtrip(self, full_summaries) -> None:
        res = build_summary(full_summaries)
        d = res.as_dict()
        assert d["workforce_total"] == 212
        assert isinstance(d["rows"], list)
        assert d["rows"][0]["department"] == "Плавильное"


class TestExportExcel:
    """Экспорт сводки в .xlsx (openpyxl)."""

    def test_creates_file(self, full_summaries, tmp_path: Path) -> None:
        res = build_summary(full_summaries)
        out = export_to_excel(res, tmp_path / "summary")
        assert out.exists()
        assert out.suffix == ".xlsx"

    def test_content_readable(self, full_summaries, tmp_path: Path) -> None:
        from openpyxl import load_workbook
        res = build_summary(full_summaries)
        out = export_to_excel(res, tmp_path / "s.xlsx")
        wb = load_workbook(out)
        ws = wb["Сводка"]
        text = "\n".join(str(c.value) for row in ws.iter_rows()
                         for c in row if c.value is not None)
        assert "Сводная ведомость" in text
        assert "ИСТ-6М1" in text
        # форматирование чисел: 58,932.0 или 58 932.0 — допускаем оба
        digits = "".join(ch for ch in text if ch.isdigit())
        assert "58932" in text.replace(",", "").replace(" ", "") \
            or digits.startswith("589320")

    def test_no_white_theme_colors(self, full_summaries,
                                   tmp_path: Path) -> None:
        """Оформление в палитре приложения: белого фона/текста нет."""
        from openpyxl import load_workbook
        res = build_summary(full_summaries)
        out = export_to_excel(res, tmp_path / "s.xlsx")
        wb = load_workbook(out)
        ws = wb["Сводка"]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5):
            for cell in row:
                rgb = getattr(cell.fill.fgColor, "rgb", None) or ""
                if cell.fill.patternType == "solid":
                    assert str(rgb)[-6:].upper() != "FFFFFF"
                color = getattr(cell.font.color, "rgb", None)
                if color is not None and str(color):
                    assert str(color)[-6:].upper() != "FFFFFF"

    def test_empty_result_still_exports(self, tmp_path: Path) -> None:
        res = SummaryResult()
        out = export_to_excel(res, tmp_path / "empty.xlsx")
        assert out.exists()
