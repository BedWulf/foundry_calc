"""Hourly/shift stress is not a fake calendar schedule."""
from types import SimpleNamespace

import pytest

from core.shift_capacity import calculate_shift_stress
from core.summary import build_summary, export_to_excel


def sample():
    return {
        "molding": {
            "batch_plan_feasible": True, "annual_molds_total": 10_000,
            "assigned_molds": 10_000, "concurrent_line_limit": 0,
            "shift_peak_lines": [
                {"line_id": "A", "metal_t_h": 2.0, "cores_h": 40},
                {"line_id": "B", "metal_t_h": 3.0, "cores_h": 20},
            ],
        },
        "melting": {"capacity_charge_t_h": 6.0, "metal_charge_t": 900,
                    "capacity_charge_t_year": 1200},
        "cores": {"capacity_cores_h": 80, "annual_cores": 5000,
                  "capacity_cores_year": 6000},
        "sand_prep": {},
        "transport": {"ladles": {"working": 3, "capacity_t": 3,
                                 "cycle_time_h": 2.0}},
    }


def test_two_lines_peak_is_not_mean_annual_output():
    r = calculate_shift_stress(sample(), shift_hours=8, ugar_pct=5)
    assert r.lines_running == 2
    liquid = r.rows[0]
    assert liquid.unit == "т/смену"
    assert liquid.demand == 40
    assert liquid.available == pytest.approx(45.6)
    assert liquid.status == "ok"
    assert r.rows[1].available == 36
    assert r.rows[1].status == "deficit"
    assert r.rows[2].demand == 480
    assert r.rows[2].status == "ok"
    assert r.rows[3].status == "unknown"
    assert r.has_deficit


def test_concurrency_limit_changes_peak_requirements():
    a = calculate_shift_stress(sample(), shift_hours=8, concurrent_lines=1,
                               ugar_pct=5)
    b = calculate_shift_stress(sample(), shift_hours=8, concurrent_lines=0,
                               ugar_pct=5)
    assert a.lines_running == 1 and b.lines_running == 2
    assert a.rows[0].demand == 24
    assert b.rows[0].demand == 40
    assert not a.has_deficit
    assert b.has_deficit


@pytest.mark.parametrize("bad", [-1, 2.5, True])
def test_invalid_concurrent_lines_rejected(bad):
    with pytest.raises(ValueError):
        calculate_shift_stress(sample(), shift_hours=8, concurrent_lines=bad)


def test_missing_pouring_mass_is_unknown_not_zero():
    s = sample()
    s["molding"]["shift_peak_lines"][1]["metal_t_h"] = None
    r = calculate_shift_stress(s, shift_hours=8, ugar_pct=5)
    assert r.rows[0].status == "unknown"
    assert r.rows[1].status == "unknown"
    assert r.rows[0].demand is None


def test_missing_ugar_does_not_confuse_charge_and_liquid_mass():
    r = calculate_shift_stress(sample(), shift_hours=8, ugar_pct=None)
    assert r.rows[0].available is None
    assert r.rows[0].status == "unknown"


def test_shift_errors_propagate_to_summary_and_excel(tmp_path):
    s = sample()
    p = SimpleNamespace(castings=[object()], annual_output_t=100,
                        shift_duration_h=8, ugar_pct=5, name="Ковши")
    res = build_summary(s, p)
    assert res.shift_has_deficit
    assert not res.ready
    assert any("СМЕННЫЙ ПИК" in warning for warning in res.warnings)
    fn = export_to_excel(res, tmp_path / "out.xlsx")
    from openpyxl import load_workbook
    book = load_workbook(fn)
    assert "Сменный сценарий" in book.sheetnames
    assert book["Сменный сценарий"]["A2"].value.startswith("Жидкий металл")


def test_no_castings_not_silent_success():
    s = sample()
    s["molding"]["shift_peak_lines"] = []
    r = calculate_shift_stress(s, shift_hours=8, ugar_pct=5)
    assert r.lines_running == 0
    assert r.rows[0].status == "unknown"
    assert not r.has_deficit
