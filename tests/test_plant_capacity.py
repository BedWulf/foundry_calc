"""Independent unit-consistent bottleneck analysis for a foundry workshop."""
from types import SimpleNamespace

import pytest

from core.plant_capacity import calculate_plant_capacity
from core.summary import build_summary, export_to_excel


def sources():
    return {
        "molding": {"annual_molds_total": 1000, "assigned_molds": 1000,
                    "batch_plan_feasible": True},
        "melting": {"metal_charge_t": 1000, "capacity_charge_t_year": 1200},
        "cores": {"annual_cores": 2000, "capacity_cores_year": 2100},
        "sand_prep": {"capacity_operations": [
            {"operation": "molding_mix", "title": "Формовочная смесь",
             "quantity": 150, "capacity_year": 200, "unit": "м³/год"}
        ]},
        "transport": {"ladles": {
            "liquid_metal_t": 900, "capacity_t": 3,
            "working": 2, "fund_h": 300, "cycle_time_h": 2,
        }},
    }


def test_annual_coverage_never_claims_schedule_verified():
    p = SimpleNamespace(castings=[object()], name="Цех")
    result = calculate_plant_capacity(sources(), p)
    assert not result.has_deficit
    assert result.unverified  # cranes and intra-shift delivery unknown
    assert all(x.status == "ok" for x in result.rows[:-1])
    assert result.rows[-1].status == "unknown"


def test_melting_demand_is_charge_not_good_casting_output():
    s = sources()
    s["melting"]["capacity_charge_t_year"] = 950
    report = calculate_plant_capacity(s, SimpleNamespace(castings=[object()]))
    row = next(r for r in report.rows if r.department == "Плавильное")
    assert row.demand == 1000
    assert row.capacity == 950
    assert row.status == "deficit"
    assert report.has_deficit


def test_independent_molding_deficit_not_hidden_by_other_departments():
    s = sources()
    s["molding"].update(assigned_molds=750, batch_plan_feasible=False)
    r = calculate_plant_capacity(s, SimpleNamespace(castings=[object()]))
    row = r.rows[0]
    assert row.status == "deficit"
    assert row.demand - row.capacity == 250


def test_missing_data_not_silently_treated_as_zero():
    s = sources()
    s["cores"].pop("capacity_cores_year")
    s["sand_prep"]["capacity_operations"][0]["capacity_year"] = None
    s["transport"]["ladles"].pop("cycle_time_h")
    r = calculate_plant_capacity(s, SimpleNamespace(castings=[object()]))
    assert not r.has_deficit
    assert all(row.status == "unknown" for row in r.rows if row.department in (
        "Стержневое", "Смесеприготовительное", "Транспорт"))


def test_mixed_sand_units_do_not_aggregate():
    s = sources()
    s["sand_prep"]["capacity_operations"].append({
        "operation": "crushing", "title": "Размельчение",
        "quantity": 300, "capacity_year": 100, "unit": "т/год"})
    r = calculate_plant_capacity(s, SimpleNamespace(castings=[object()]))
    sand = [x for x in r.rows if x.department == "Смесеприготовительное"]
    assert len(sand) == 2
    assert sand[0].unit == "м³/год"
    assert sand[1].unit == "т/год"
    assert sand[1].status == "deficit"


@pytest.mark.parametrize("bad", [None, float("nan"), float("inf"), -1, "n/a"])
def test_invalid_capacity_does_not_become_numeric_zero(bad):
    s = sources()
    s["melting"]["capacity_charge_t_year"] = bad
    row = next(x for x in calculate_plant_capacity(s).rows if x.department == "Плавильное")
    assert row.status == "unknown"
    assert row.capacity is None


def test_summary_flags_melting_deficit_and_exports_comparable_units(tmp_path):
    s = sources()
    s["melting"]["capacity_charge_t_year"] = 500
    p = SimpleNamespace(castings=[object()], name="Баланс", annual_output_t=100)
    res = build_summary(s, p)
    assert not res.ready
    assert res.capacity_has_deficit
    assert any("НЕДОСТАТОК МОЩНОСТИ" in x for x in res.warnings)
    assert "capacity_rows" in res.as_dict()
    filename = export_to_excel(res, tmp_path / "capacity.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(filename)
    assert "Баланс мощностей" in wb.sheetnames
    sheet = wb["Баланс мощностей"]
    rows = list(sheet.values)
    assert any(row[1] == "Металлозавалка" and row[2] == 1000
               and row[3] == 500 for row in rows[1:])


def test_sand_operation_incompatible_declared_units_are_unknown():
    s = sources()
    s["sand_prep"]["capacity_operations"][0].update(
        quantity=150, capacity_year=200,
        quantity_unit="м³/год", capacity_unit="т/год")
    report = calculate_plant_capacity(s, SimpleNamespace(castings=[object()]))
    row = next(x for x in report.rows if x.department == "Смесеприготовительное")
    assert row.status == "unknown"
    assert row.demand == 150
    assert row.capacity is None
    assert "Несовместимые единицы" in row.note


def test_sand_operation_matching_declared_units_still_compares():
    s = sources()
    s["sand_prep"]["capacity_operations"][0].update(
        quantity_unit="м³/год", capacity_unit="м³/год")
    report = calculate_plant_capacity(s)
    row = next(x for x in report.rows if x.department == "Смесеприготовительное")
    assert row.status == "ok"
    assert row.capacity == 200


def test_sand_explicit_volume_demand_vs_mass_capacity_marked_unknown():
    s = sources()
    s["sand_prep"]["capacity_operations"][0].update(
        quantity_unit="м³", capacity_unit="т/год")
    row = next(x for x in calculate_plant_capacity(s).rows
               if x.department == "Смесеприготовительное")
    assert row.status == "unknown"


def test_sand_demand_unit_checked_against_legacy_row_unit():
    s = sources()
    s["sand_prep"]["capacity_operations"][0]["quantity_unit"] = "т/год"
    row = next(x for x in calculate_plant_capacity(s).rows
               if x.department == "Смесеприготовительное")
    assert row.status == "unknown"


def test_sand_volume_without_annual_suffix_matches_annual_volume():
    s = sources()
    s["sand_prep"]["capacity_operations"][0]["quantity_unit"] = "м³"
    row = next(x for x in calculate_plant_capacity(s).rows
               if x.department == "Смесеприготовительное")
    assert row.status == "ok"
