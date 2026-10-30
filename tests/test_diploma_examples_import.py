"""Regression tests for both REAL committed diploma XLSX workbooks.

They are source-derived test fixtures, not fabricated in the test itself.
Original numerical discrepancies must stay in their source sheets.
"""
from pathlib import Path
import math
import pytest
from openpyxl import load_workbook

from import_export.excel_importer import ExcelImporter

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = [
    ("Диплом_1_ведомость_формовки.xlsx", 5_271_000, 1_368_325),
    ("Диплом_2_ведомость_формовки.xlsx", 9_001_281, 768_159),
]


@pytest.mark.parametrize("filename,annual_pieces,physical_molds", SAMPLES)
def test_original_diploma_workbook_is_importable_and_reconciled(
    filename, annual_pieces, physical_molds
):
    path = ROOT / "examples" / filename
    assert path.is_file(), f"Отсутствует контрольная ведомость: {path}"
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        assert wb.sheetnames == ["Ведомость формовки", "Импорт A-P", "Сверка"]
        ws = wb.worksheets[0]
        source_rows = list(ws.iter_rows(min_row=4, max_row=23, values_only=True))
        assert len(source_rows) == 20
        assert sum(int(row[5]) for row in source_rows) == annual_pieces
        assert sum(math.ceil(int(row[5]) / int(row[6])) for row in source_rows) == physical_molds
        assert ws["A24"].value == "ИТОГО (PDF)"
    finally:
        wb.close()

    preview = ExcelImporter().load_preview(path)
    assert len(preview.castings) == 20
    assert sum(x.annual_qty_with_brak for x in preview.castings) == annual_pieces
    assert sum(x.molds_per_year for x in preview.castings) == physical_molds
    assert all(x.molds_per_year > 0 and x.pouring_mass_kg > 0
               and x.annual_metal_t > 0 for x in preview.castings)
    assert all(x.flask_size_mm.startswith("900x600x")
               for x in preview.castings)
    assert len(preview.warnings) > 0  # keep source disagreements visible


@pytest.mark.parametrize("filename,annual_pieces,physical_molds", SAMPLES)
def test_import_a_p_sheet_matches_unmodified_diploma_source(
    filename, annual_pieces, physical_molds, tmp_path
):
    source = ROOT / "examples" / filename
    wb = load_workbook(source)
    try:
        # Import A:P was generated for ordinary Excel import. Make it the first
        # worksheet to simulate a student saving only the normalized table.
        worksheet = wb["Импорт A-P"]
        wb._sheets.remove(worksheet)
        wb._sheets.insert(0, worksheet)
        copy = tmp_path / filename
        wb.save(copy)
    finally:
        wb.close()
    converted = ExcelImporter().load_preview(copy)
    assert len(converted.castings) == 20
    assert sum(c.annual_qty_with_brak for c in converted.castings) == annual_pieces
    assert sum(c.molds_per_year for c in converted.castings) == physical_molds
