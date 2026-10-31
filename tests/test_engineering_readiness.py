"""Final-report readiness cannot silently trust example steel/iron material data."""
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from core.project import Project
from core.engineering_readiness import audit_calculation_sources, REQUIRED_MATERIAL_FIELDS
from core.summary import build_summary, export_to_excel


def test_legacy_steel_has_five_undocumented_physical_properties():
    p = Project(selected_alloy_id="steel_20L")
    findings = audit_calculation_sources(p)
    assert len(findings) == len(REQUIRED_MATERIAL_FIELDS)
    for field in REQUIRED_MATERIAL_FIELDS:
        assert any(field in issue for issue in findings)


def test_cast_iron_has_same_fail_closed_policy():
    p = Project(selected_alloy_id="cast_iron_SCh20")
    problems = audit_calculation_sources(p)
    assert len(problems) == 5
    assert all("СЧ20" in item for item in problems)


def test_reference_only_indotherm_does_not_certify_furnace_count():
    p = Project(selected_alloy_id="steel_20L")
    sources = audit_calculation_sources(p, {
        "furnace_name": "Indotherm casting unit 1000 kg / 500 kW",
        "model_designation": "INDOTHERM-CU-1000-500",
        # Published reference-only melting rate is NOT an approved
        # field-level production parameter for the specific plant.
    })
    assert any("Печь" in p and "подтверждённых" in p for p in sources)


def test_only_complete_grade_and_field_provenance_is_accepted(tmp_path):
    directory = tmp_path / "materials"
    directory.mkdir()
    grade = "20Л"
    entries = {"id": "steel_20L", "name": grade, "category": "steel"}
    for name in REQUIRED_MATERIAL_FIELDS:
        entries[name] = 1.0 if name != "burnout_pct" else {"C": 1}
    entries["field_provenance"] = {
        field: {
            "grade": grade, "field": field,
            "verification_status": "verified",
            "source_type": "technical_reference",
            "document_url": "https://example.org/source",
            "document_title": "Source document",
            "document_revision": "2026",
            "location": "Table 1",
            "measurement_conditions": "Test conditions",
        }
        for field in REQUIRED_MATERIAL_FIELDS
    }
    (directory / "alloys.json").write_text(
        json.dumps({"alloys": [entries]}, ensure_ascii=False), encoding="utf-8"
    )
    p = SimpleNamespace(selected_alloy_id="steel_20L")
    # Self-declared "verified" must NEVER grant release approval.
    (tmp_path / "verified_references").mkdir()
    registry_path = tmp_path / "verified_references" / "approved_operational_material_properties.json"
    registry = {
        "usage_policy": "fail_closed_no_approvals_until_documented",
        "approved_records": [],
    }
    registry_path.write_text(json.dumps(registry), encoding="utf-8")
    assert len(audit_calculation_sources(p, data_dir=tmp_path)) == 5
    # Even after adding manifest records, the exact values and sources
    # must match; this only tests matching logic, NOT source authenticity.
    registry["approved_records"] = [
        {"alloy_id": "steel_20L", "field": field,
         "value": entries[field], "source": entries["field_provenance"][field],
         "review_status": "approved", "reviewer": "test-reviewer",
         "review_date": "2026-10-31"}
        for field in REQUIRED_MATERIAL_FIELDS
    ]
    registry_path.write_text(json.dumps(registry), encoding="utf-8")
    assert audit_calculation_sources(p, data_dir=tmp_path) == []
    entries["field_provenance"]["density_kg_m3"]["grade"] = "35Л"
    (directory / "alloys.json").write_text(
        json.dumps({"alloys": [entries]}, ensure_ascii=False), encoding="utf-8"
    )
    errors = audit_calculation_sources(p, data_dir=tmp_path)
    assert len(errors) == 1 and "density_kg_m3" in errors[0]


def test_project_summary_never_marks_legacy_steel_as_ready():
    project = Project(selected_alloy_id="steel_20L")
    result = build_summary({
        "melting": {
            "metal_charge_t": 100,
            "furnace_name": "ИСТ-6М1",
            "furnace_count": 1,
            "load_factor": .75,
        }
    }, project)
    assert not result.ready
    assert len(result.evidence_blockers) >= 5
    assert any("НЕ ПОДТВЕРЖДЕНЫ" in x for x in result.warnings)
    assert result.as_dict()["evidence_blockers"]


def test_export_contains_visible_source_check_sheet(tmp_path):
    result = build_summary({
        "melting": {
            "metal_charge_t": 100,
            "furnace_name": "ИСТ-6М1",
            "furnace_count": 1,
            "load_factor": .75,
        }
    }, Project(selected_alloy_id="steel_20L"))
    p = export_to_excel(result, tmp_path / "preliminary.xlsx")
    workbook = load_workbook(p, read_only=True)
    assert "Проверка источников" in workbook.sheetnames
    ws = workbook["Проверка источников"]
    text = "\n".join(
        str(cell.value) for row in ws.iter_rows()
        for cell in row if cell.value is not None
    )
    assert "БЛОКИРУЕТ ГОТОВНОСТЬ" in text
    assert "density_kg_m3" in text
    assert "ИСТ-6М1" in text
    assert "НЕПОЛНЫЙ ПРОЕКТ" in str(workbook["Сводка"]["A3"].value)
    workbook.close()


def test_missing_material_approval_manifest_fails_closed(tmp_path):
    (tmp_path / "materials").mkdir()
    (tmp_path / "materials" / "alloys.json").write_text(
        json.dumps({"alloys": [{"id": "steel_20L", "name": "20Л",
                                 "category": "steel", **{
                                     field: 1 for field in REQUIRED_MATERIAL_FIELDS
                                 }}]}, ensure_ascii=False), encoding="utf-8"
    )
    p = Project(selected_alloy_id="steel_20L")
    # A complete alloy card without the independent approval manifest fails.
    assert "Реестр допуска" in audit_calculation_sources(
        p, data_dir=tmp_path
    )[0]


def test_legacy_melting_furnace_is_not_approved_by_normative_grade():
    res = build_summary({
        "melting": {"furnace_name": "ИСТ-6М1", "furnace_count": 6,
                    "metal_charge_t": 58932.0, "load_factor": .84}
    }, Project(selected_alloy_id="steel_20L"))
    assert res.ready is False
    assert any("Печь ИСТ-6М1" in s for s in res.evidence_blockers)
