"""Cross-check original teacher course workbook without changing its values."""
from openpyxl import Workbook, load_workbook
from import_export.excel_importer import ExcelImporter


def teacher_file(tmp_path):
    wb = Workbook()
    t1 = wb.active
    t1.title = "Таблица 1"
    t1["A6"] = "Корпус"
    t1["B6"] = 1
    t1["I6"] = 8.0
    t2 = wb.create_sheet("Таблица 2")
    t2["A8"] = "Корпус"
    t2["B8"] = 1
    t2["O8"] = 100.2
    t8 = wb.create_sheet("Таблица 8")
    t8["A1"] = "Номер"
    t8["B1"] = "Наименование отливок"
    t8["G1"] = "Количество отливок с учетом брака, шт."
    t8["H1"] = "Количество отливок в форме"
    source_row = [1, "Корпус", 8, 120, 80, 50, 100, 2, 50.0,
                  "900х600", "125/125", 18, 900, .12, 6, "ИЛ-225"]
    for col, val in enumerate(source_row, 1):
        t8.cell(7, col, val)
    target = tmp_path / "teacher.xlsx"
    wb.save(target)
    return target


def test_agreeing_tables_have_no_crosscheck_errors(tmp_path):
    path = teacher_file(tmp_path)
    preview = ExcelImporter().load_preview(path)
    assert len(preview.castings) == 1
    assert not any("Таблице 1 имя" in w for w in preview.warnings)
    assert not any("не согласуется с Таблицей 8" in w for w in preview.warnings)


def test_mismatch_is_visible_without_rewriting_teacher_values(tmp_path):
    path = teacher_file(tmp_path)
    wb = load_workbook(path)
    wb["Таблица 1"]["I6"] = 99
    wb["Таблица 2"]["O8"] = 250
    wb.save(path)
    preview = ExcelImporter().load_preview(path)
    assert len(preview.castings) == 1
    assert preview.castings[0].mass_with_LPS_kg == 8
    assert preview.castings[0].annual_qty_with_brak == 100
    assert any("Таблице 1 имя или" in msg for msg in preview.warnings)
    assert any("не согласуется с Таблицей 8" in msg for msg in preview.warnings)
