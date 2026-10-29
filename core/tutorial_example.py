"""Небольшой безопасный образец Excel для обучения (не изменяет проект)."""
from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def save_tutorial_excel(filename: Path) -> Path:
    """Создать пример, который читает штатный ExcelImporter (A:P, строка 4+).

    Колонка F — отливок в год С УЧЁТОМ брака;
    G — отливок на форму; L — металл в одну форму.
    Числа учебные, а не нормативное обоснование оборудования.
    """
    target = Path(filename)
    if target.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Файл учебной номенклатуры должен иметь расширение .xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Учебная номенклатура"
    ws.append(["Учебный пример для ЛитейныйЦех PRO — не производственный проект"])
    ws.append(["F = отливок в год с учётом брака; G = отливок в форме; "
               "L = кг заливаемого металла на форму"])
    headers = [
        "№", "Отливка", "Масса с ЛПС, кг", "Габариты, мм",
        "Высота, мм", "Отливок/год с браком", "Отливок/форма",
        "Расч. металл, т", "Опока L, мм", "Опока W, мм",
        "Высоты половин, мм", "Заливочная масса, кг/форма",
        "Расч. заливка, т", "Смесь, м³", "Примечание", "Тип машины",
    ]
    ws.append(headers)
    rows = [
        (1, "Учебный корпус", 8.0, "150x120", 100, 1200, 2,
         None, 900, 600, "200/200", 18.0, None, 0.18, None, "ИЛ-225"),
        (2, "Учебная крышка", 5.0, "180x140", 80, 900, 3,
         None, 900, 600, "200/200", 17.0, None, 0.18, None, "ИЛ-225"),
    ]
    for row in rows:
        ws.append(row)
    for col in range(1, 17):
        c = ws.cell(row=3, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="36516E")
        c.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[c.column_letter].width = (31 if col in (2, 6, 12) else 18)
    ws.row_dimensions[3].height = 46
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = "A3:P5"
    wb.save(target)
    return target
